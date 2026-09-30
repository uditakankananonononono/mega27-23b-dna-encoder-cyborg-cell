"""Idempotently append audited baseline section to the corrected paper artifact.
Does not regenerate legacy make_paper50.py claims. Input is reviewed corrected DOCX.
"""
import json,sys
from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'experiments'))
import paper50 as P
path=ROOT/'paper/MEGA27-23b-50p.docx';doc=Document(path)
TITLE='External Fountain port and strand-geometry control'
# Remove only this optional section and everything following it on reruns.
ps=list(doc.paragraphs);hit=next((p for p in ps if p.text==TITLE),None)
if hit is not None:
 node=hit._element
 while node is not None:
  nxt=node.getnext()
  if node.tag.endswith('sectPr'):break
  node.getparent().remove(node);node=nxt
f=json.loads((ROOT/'results/fountain_port_audit.json').read_text());s=json.loads((ROOT/'results/rsns_segmented_audit.json').read_text())
P.h1(doc,TITLE)
P.h2(doc,'Pinned implementation and clean-channel verification')
P.para(doc,'After the local CRC8/uniform-degree baseline failed clean recovery, we tested the separately compiled, unmodified Python-3 port at https://github.com/jdbrody/dna-fountain, commit '+f['source_commit']+'. The upstream repository is https://github.com/TeamErlich/dna-fountain. This is a software-port audit, not a reproduction of experimental sequencing. Modern reedsolo returns a tuple; the port correctly takes decode(data)[0]. A return-type failure did not explain the observed noisy recovery. Clean decoded files were checked by byte equality and stored SHA-256 identity, not merely process exit.')
P.para(doc,'The pilot uses four independently generated 2,048-byte messages with payload seed 511. Each message has three separate noise realizations per rate. Twelve trials are therefore four message identities times three noise repeats, not twelve independent messages. At zero noise the three repeats duplicate the clean case. The external encoder uses 32-byte chunks, two RS parity bytes, robust-soliton delta 0.001 and c 0.025, homopolymer maximum 3 and GC fraction 0.45-0.55. Its fixed budget is 180 oligos of 152 nucleotides each. Both the intact local RSNS and the external port recover all four clean messages.')
P.h2(doc,'Fixed-budget strict-substitution results')
rows=[]
for rate in [0,.005,.01,.02,.03]:
 vals=[]
 for codec in ['external_fountain_py3_port','repaired_rsns']:
  rr=[r for r in f['rows'] if r['codec']==codec and r['p']==rate];vals.append(str(sum(r['recovered'] for r in rr))+'/12')
 rr=[r for r in s['rows'] if r['p']==rate];vals+=[str(sum(r['whole_file_recovered'] for r in rr))+'/12',str(sum(r['chunks_recovered'] for r in rr))+'/768']
 rows.append([str(100*rate)+'%']+vals)
P.table(doc,'Table. Whole-file recovery and segmented local chunk recovery. Trials share four payload identities; the chunk total is not an independent-message sample size.',['Changed-base probability','External port files','Intact RSNS files','Segmented RSNS files','Segmented chunks'],rows)
P.para(doc,'Noise is independent Bernoulli selection followed by a uniform choice among the three other bases. Requested probabilities are changed-base probabilities, unlike the old four-choice attempt channel. Unequal strand lengths consume noise random numbers differently, so these are not matched identical error locations. Every realized fraction, payload hash and trial outcome is in results/fountain_port_audit.json and results/rsns_segmented_audit.json. The segmented run retains every chunk success and decoder exception. No old trial was replaced.')
P.h2(doc,'Strand geometry changes the density statement')
P.para(doc,'The intact local RSNS sequence has 18,495 nucleotides. It is not interchangeable with the external port\'s 152-nucleotide oligos. To isolate one practical consequence, we split the identical payloads into 64 independent 32-byte RSNS pieces, each 495 nucleotides long. This control assumes perfect known order and chunk identity, charges no index bases, and has no outer redundancy. It is not yet a physically matched storage design: its oligos remain more than three times longer than the external oligos.')
P.table(doc,'Table. Encoded-base accounting, excluding primers and molecular-copy/read budgets.',['Design','Strands','Nucleotides/strand','Total encoded nt','Payload bits/encoded nt'],[['External port','180','152','27,360','0.598830'],['Intact RSNS','1','18,495','18,495','0.885861'],['Segmented RSNS','64','495','31,680','0.517172']])
P.para(doc,'Thus the local density ordering reverses when the same codec is independently framed in short pieces: intact RSNS is denser than this fixed port budget, whereas segmented RSNS is less dense. Repeated headers and rounding each piece to complete RS blocks explain this cost without invoking a biological mechanism. At 1% the segmented design loses five of 768 chunks, yet five failures can cause five entire files to fail; it recovers seven of twelve files. At 2%, 708 of 768 chunks survive but no complete file survives. Reporting only chunk recovery would hide this file-level penalty. These denominators measure different outcomes.')
P.h2(doc,'What the pilot establishes and what remains open')
P.para(doc,'This pilot replaces an invalid local baseline with a pinned external software implementation that passes clean-channel recovery, and demonstrates a resource-accounting sensitivity within RSNS. It does not establish a fair win over DNA Fountain. Both encoder budgets are fixed and unoptimized, the local segmented design has ideal uncharged indexing, the external port has only two parity bytes per droplet, and no read consensus is used. There are no primers, synthesis constraints, molecule-copy budgets, strand dropout, insertion/deletion or measured sequencing errors in this test. A balanced comparison must match permitted strand lengths, total encoded bases and read budgets, account for ordering metadata, and sweep redundancy for both methods before making a frontier claim.')
P.para(doc,'Only four independent messages support this pilot. Noise repeats cannot justify a narrow population confidence interval or a general message-distribution claim. The external port\'s decline under this one-read strict channel is an observed software result at these parameters, not a contradiction of experimental storage demonstrations. Published-codec superiority and biological validation remain unresolved. Reproduction: python3 experiments/fountain_port_audit.py --port /path/to/pinned/compiled/port; python3 experiments/rsns_segmented_audit.py; python3 experiments/append_fountain_audit.py. Dependencies and provenance are retained in the external audit JSON.')
from docx.oxml import OxmlElement
for t in doc.tables[-2:]:
 for row in t.rows:
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 for cell in t.rows[0].cells:
  for paragraph in cell.paragraphs:
   paragraph.paragraph_format.keep_with_next=True
P.save(doc,str(path));print('appended audited external section')
