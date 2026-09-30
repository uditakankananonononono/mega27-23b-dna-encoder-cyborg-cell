"""Append exact clean graph result to corrected existing paper; retain history."""
import json,sys
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'experiments'));import paper50 as P
path=ROOT/'paper/MEGA27-23b-50p.docx';d=Document(path);title='Clean-channel threshold: screening and peeling at near-matched geometry'
hit=next((p for p in d.paragraphs if p.text==title),None)
if hit is not None:
 node=hit._element
 while node is not None:
  nxt=node.getnext()
  if node.tag.endswith('sectPr'):break
  node.getparent().remove(node);node=nxt
j=json.loads((ROOT/'results/fountain_clean_graph_audit.json').read_text());rows=j['rows']
P.h1(d,title)
P.para(d,'The next comparison kept external oligo geometry close to the independently framed local control: 63 oligos of 496 nucleotides, totaling 31,248 bases, against local 64 by 495 nucleotides, totaling 31,680 bases. The external design uses 1.36% fewer encoded bases and carries seed headers. Local ordering remains ideal and uncharged. Payloads are the same four 2,048-byte messages. We swept two, eight, sixteen and twenty-four external RS parity bytes while changing payload bytes per droplet to 118, 112, 104 and 96, respectively. This holds encoded oligo length fixed. Known source length removes terminal zero padding in both evaluations. A failed clean-channel gate stopped noisy comparative interpretation.')
P.para(d,'At maximum homopolymer length three and GC fraction 0.45-0.55, the port peeling decoder fails all four clean messages at every parity setting. Clean failures cannot be interpreted as poor substitution correction. We therefore used an exact graph diagnostic: regenerate source-chunk incidence from embedded seeds, eliminate the binary incidence matrix over GF(2), and propagate the payload right-hand sides. Full-rank systems are back-substituted and checked against original bytes. This is a clean mathematical oracle, not a modified port decoder or a noisy-error-correcting method. The port source remains unmodified.')
P.table(d,'Table. Clean recovery at fixed 63 by 496-nt budget. Each count is out of four distinct payloads; graph solve is an oracle, not the original decoding algorithm.',['RS bytes','HP 3 peel / linear','HP 4 peel / linear','HP 5 peel / linear'],[[str(rs)]+[str(sum(r['port_clean_recovered'] for r in rows if r['rs_bytes']==rs and r['max_homopolymer']==hp))+'/4 ; '+str(sum(r['clean_linear_oracle_recovered'] for r in rows if r['rs_bytes']==rs and r['max_homopolymer']==hp))+'/4' for hp in [3,4,5]] for rs in [2,8,16,24]])
P.h2(d,'Mechanism within the software graph')
P.para(d,'The two- and eight-byte parity configurations have zero accepted degree-one oligos for every payload under the HP 3 screen. Peeling cannot start without a singleton, even when the incidence matrix has full rank. At RS 24, all four clean matrices have full rank and exact elimination recovers all messages, while peeling recovers none. At RS 16, only two of four matrices have full rank. This separates an algorithmic stopping-set problem from missing information. Many degree-one candidates can be copies of the same source chunk; oligo count alone does not measure coverage of distinct singleton chunks. Every degree histogram, distinct singleton count, binary rank, accepted-sequence hash, screen-attempt count and decoded-file outcome is retained.')
P.para(d,'Relaxing the screen to maximum homopolymer four produces clean peeling 4/4, 3/4, 4/4 and 4/4 across the parity grid. The remaining RS 8 failure has a full-rank matrix but zero singleton chunks. At maximum homopolymer five, all configurations recover all four clean messages with the original peeling decoder. Total oligo count, length, GC window, robust-soliton c=0.025 and delta=0.001 remain fixed. These threshold observations support a screening/graph interaction in this small-payload regime. They do not establish the performance of an optimized Fountain design, because the degree distribution, seed stream and accepted budget were not optimized.')
P.h2(d,'Limits and next valid comparison')
P.para(d,'Changing permitted homopolymer length changes the sequence constraint. Clean gains from relaxing it are not free gains under an unchanged synthesis requirement. Exact linear recovery also changes the algorithm and applies only to validated clean data here. A noisy comparison may use configurations that pass the clean gate, but must name their distinct constraints and account for local ideal indexing. Four payloads and the port default deterministic LFSR stream do not establish general seed-distribution performance. No biological synthesis, consensus reads, molecule-copy counts, primers, strand dropout or indels were tested. A fair frontier remains open.')
P.para(d,'The initial HP 3 strict-noise sweep on message zero is preserved separately in results/fountain_geometry_message_0.json and its clean encoded FASTA files. Its failures are not benchmark evidence against the published method. The complete four-message clean audit is results/fountain_clean_graph_audit.json, with per-message files for bounded reproduction. Run python3 experiments/fountain_clean_graph_audit.py --port /path/to/pinned/compiled/port --message N for N=0,1,2,3. The pinned external source is https://github.com/jdbrody/dna-fountain at f97c1b8a81f5c5b819209d5b5e26c9c8f4439495. Published-codec superiority and biological validation remain withdrawn.')
for row in d.tables[-1].rows:
 row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
for cell in d.tables[-1].rows[0].cells:
 for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
P.save(d,str(path));print('appended clean graph audit')
