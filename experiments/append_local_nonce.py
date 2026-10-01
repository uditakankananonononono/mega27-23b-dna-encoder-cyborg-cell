from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
def main():
 p=ROOT/'paper/MEGA27-23b-50p.docx';d=Document(p);title='A nonce-only substitution can bypass core payload identity'
 assert not any(x.text==title for x in d.paragraphs)
 d.add_heading(title,level=1)
 d.add_paragraph('The prospective codec exposes a failure surface outside the isolated 48-nucleotide theorem. A frozen diagnostic exhausts all four positions and three alternative letters of the first data-block nonce in each of the same eight clean identities. Each mutation is applied identically to all three copies while the 48-base core remains untouched. There are 288 events across three detectors, not 288 independent files or a random channel simulation. We separately classify core acceptance after unmasking and whole-file decoding.')
 t=d.add_table(rows=1,cols=5);t.style='Light Shading Accent 1'
 for c,v in zip(t.rows[0].cells,['Detector','Spacing reject','Core reject','Core wrong accept','File exact/loud']):c.text=v
 for det in ['Low check','Full check','Distance five']:
  for c,v in zip(t.add_row().cells,[det,'56/96','16/96','24/96','72/96; 24/96']):c.text=v
 d.add_paragraph('For each detector, 24 of 96 nonce mutations are accepted by the core but produce a wrong unmasked payload. These are prefix positions zero through two whose accepted nonce chooses a different mask while leaving the core transition seed unchanged. A zero core syndrome then authenticates the masked codeword, not its association with the intended nonce. This mechanism is not a violation of the core distance proof: the protected core did not change. It is a boundary on what that proof can say about a framed file.')
 d.add_paragraph('Whole-file decoding returns the exact original in 72 cases per detector because a rejected first block is a single erasure recoverable with the saved outer parity. The 24 accepted-but-wrong cases instead fail loudly under the global header/checksum/byte checks. No silent wrong is observed in these 288 deliberately correlated events. That containment does not prove a never-wrong guarantee, cover a different data or parity block, or cover multiple nonce/core edits, indels, dropout or physical errors. Three copies do not repair a mutation deliberately shared by all three.')
 d.add_paragraph('The result requires separate claims for block-code detection, framing integrity, erasure recovery and final-file validation. A nonce-aware protected frame or stronger global integrity design would be a different codec and must charge its overhead and test its own proof boundary. We do not silently reinterpret this audit as a production repair. The frozen plan, complete event ledger and regression are in experiments/local_nonce_audit_plan.json, results/local_nonce_audit.json and tests/test_local_nonce_audit.py. All earlier long-file failures and local-constraint caveats remain visible.')
 d.save(p)
if __name__=='__main__':main()
