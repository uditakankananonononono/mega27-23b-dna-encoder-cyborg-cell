"""Append source qualification to the current manuscript without losing prior audits."""
from pathlib import Path
import json
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
p=ROOT/'paper/MEGA27-23b-50p.docx';d=Document(p);section=json.loads((ROOT/'research/dna_fountain_source_paper_section.json').read_text())
if section[0] not in [r.text for r in d.paragraphs]:
 d.add_page_break();d.add_heading(section[0],level=1)
 for paragraph in section[1:]:d.add_paragraph(paragraph)
 d.save(p)
