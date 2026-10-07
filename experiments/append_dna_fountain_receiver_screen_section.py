"""Preserve prior manuscript while appending the frozen receiver-screen report."""
from pathlib import Path
import json
from docx import Document
ROOT=Path(__file__).resolve().parents[1];p=ROOT/'paper/MEGA27-23b-50p.docx';d=Document(p);s=json.loads((ROOT/'research/dna_fountain_receiver_screen_section.json').read_text())
if s[0] not in [r.text for r in d.paragraphs]:
 d.add_heading(s[0],level=1)
 for t in s[1:]:d.add_paragraph(t)
 d.save(p)
