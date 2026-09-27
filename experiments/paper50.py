"""paper50: thin python-docx wrapper for 50-page paper generators.
Rebuilt 2026-09-27 (original lost with prior sandbox). Times New Roman throughout
(user formatting rule). API: new_doc, title_block, h1, h2, h3, para, eq, table,
figure, page_break, save.
"""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

FONT = "Times New Roman"

def _style_run(r, size=12, bold=False, italic=False):
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic

def new_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    return doc

def title_block(doc, title, subtitle):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(p.add_run(title), 17, bold=True)
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(p2.add_run(subtitle), 12, italic=True)
    doc.add_paragraph()

def h1(doc, text):
    p = doc.add_heading(level=1); r = p.add_run(text); _style_run(r, 15, bold=True)

def h2(doc, text):
    p = doc.add_heading(level=2); r = p.add_run(text); _style_run(r, 13, bold=True)

def h3(doc, text):
    p = doc.add_heading(level=3); r = p.add_run(text); _style_run(r, 12, bold=True)

def para(doc, text):
    p = doc.add_paragraph()
    _style_run(p.add_run(text), 12)

def eq(doc, num, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(p.add_run(f"({num})   {text}"), 12, italic=True)

def table(doc, caption, headers, rows):
    para(doc, caption)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ""
        _style_run(cell.paragraphs[0].add_run(str(h)), 10, bold=True)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            cell.text = ""
            _style_run(cell.paragraphs[0].add_run(str(v)), 10)
    doc.add_paragraph()

def figure(doc, path, caption):
    try:
        doc.add_picture(path, width=Inches(6.0))
    except Exception as e:
        para(doc, f"[figure unavailable: {path} ({e})]")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(p.add_run(caption), 10, italic=True)

def page_break(doc):
    doc.add_page_break()

def save(doc, path):
    doc.save(path)
