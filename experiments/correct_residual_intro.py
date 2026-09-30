"""Residual whole-manuscript claim correction with verbatim audit ledger."""
import json,sys
from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parents[1];p=ROOT/'paper/MEGA27-23b-50p.docx';d=Document(p)
changes={
14:'Part I tests a local biologically constrained codec family using independently seeded payload blocks, incomplete local error detection and one conditional erasure-repair block. Substitutions in the inverse rotating-base mapping affect adjacent trits rather than a downstream suffix. The current low-byte check does not turn every corruption into a known erasure; a saved strict-noise file also reproduces a silent wrong output in the unmodified production decoder under a stated tie seed. Historical local baseline trials and later external-port audits are retained separately. No published-baseline superiority, universal integrity or physical storage validation is established.',
66:'The retained evolution_error_models.py proxy experiment reports 16/16 recovery under its Illumina-like parameters, 1/16 under its synthesis-like parameters and 0/16 under its Nanopore-like parameters for the tested local RS configuration. These are finite software observations at assigned substitution/indel and homopolymer-multiplier settings, not calibrated instrument measurements or a guarantee that the code covers a named read technology. They neither measure the redundancy needed for an actual laboratory channel nor certify a physical storage design. The source experiment, parameters and negative outcomes remain retained; biological validation is open.',
165:'The codec experiments model chosen software sequence constraints, not manufacturing success. Continuous never-same encoding excludes adjacent repeated bases, whereas independent seeded blocks can join with two-base repeats. Scrambling gives measured GC distributions but no universal sliding-window bound. Synthesis yield, sequencing error, amplification and vendor acceptance depend on protocols and sequence contexts that were not measured here. Earlier numerical process efficiencies and technology-specific run limits in this paragraph were unsupported by linked primary evidence and are withdrawn. No claim that this implementation removes a whole physical failure class is established.'
}
ledger=[]
for i,new in changes.items():
 old=d.paragraphs[i].text;ledger.append({'paragraph_index':i,'before':old,'after':new});d.paragraphs[i].text=new
 # Preserve house-style font after replacing runs.
 for run in d.paragraphs[i].runs:run.font.name='Times New Roman'
for i,x in enumerate(d.paragraphs):
 if x.text=='9j. Realistic error profiles: the codec covers Illumina-scale channels exactly':
  old=x.text;x.text='9j. Assigned error-profile proxies: finite recovery without instrument validation';ledger.append({'paragraph_index':i,'before':old,'after':x.text})
d.save(p);(ROOT/'results/residual_intro_claim_corrections.json').write_text(json.dumps(ledger,indent=2)+'\n')
