"""Append clean-valid noisy parity sweep without regenerating legacy claims."""
import json,sys
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'experiments'));import paper50 as P
path=ROOT/'paper/MEGA27-23b-50p.docx';d=Document(path);title='Clean-valid noisy comparison: parity allocation reverses the pilot interpretation'
hit=next((p for p in d.paragraphs if p.text==title),None)
if hit is not None:
 node=hit._element
 while node is not None:
  nxt=node.getnext()
  if node.tag.endswith('sectPr'):break
  node.getparent().remove(node);node=nxt
j=json.loads((ROOT/'results/fountain_valid_noise_audit.json').read_text());s=json.loads((ROOT/'results/rsns_segmented_audit.json').read_text());rows=j['rows']
P.h1(d,title)
P.para(d,'We next used only external configurations that passed the clean-channel gate on every tested payload. The port permits homopolymers up to length five, retains GC fraction 0.45-0.55, and produces 63 oligos of 496 nucleotides. Four parity settings preserve the encoded-base budget of 31,248 nucleotides. These settings allocate two, eight, sixteen or twenty-four bytes to inner RS parity and correspondingly 118, 112, 104 or 96 bytes to source chunks. The same four 2,048-byte payload identities and their hashes match the retained segmented RSNS control. All zero-noise byte-equality checks pass. Original port code is unmodified, and its original peeling decoder is used, not the clean Gaussian oracle.')
P.para(d,'For each message and rate, three independent strict-substitution noise seeds are applied. Each encoded oligo receives one read and no consensus. Changed-base probabilities are zero, 0.5%, 1%, 2% and 3%, with a uniform choice among the other three bases after each Bernoulli selection. Zero-noise repeats duplicate a clean input; the denominator twelve is four message identities times three repetitions, not twelve independent messages. The local and external strands are almost the same length but different sequences, so equal random-number seeds do not make identical physical error locations. All realized fractions and exceptions are retained in results/fountain_valid_noise_audit.json and per-message JSON files.')
out=[]
for p in [0,.005,.01,.02,.03]:
 a=[str(100*p)+'%']
 for rs in [2,8,16,24]:a.append(str(sum(r['recovered'] for r in rows if r['rs_bytes']==rs and r['p']==p))+'/12')
 a.append(str(sum(r['whole_file_recovered'] for r in s['rows'] if r['p']==p))+'/12');out.append(a)
P.table(d,'Table. Whole-file success. External HP5 versus local HP1, with near-matched lengths and encoded budget, not identical constraints.',['Changed-base p','Port RS2','Port RS8','Port RS16','Port RS24','Segmented RSNS'],out)
P.h2(d,'Redundancy allocation is an experimental variable')
P.para(d,'At 1%, changing external inner parity from two to sixteen bytes changes complete-file recovery from zero to twelve of twelve within the same total encoded-base budget. At 2%, the twenty-four-byte setting recovers twelve of twelve while the sixteen-byte setting recovers one of twelve and the segmented local control recovers zero. At 3%, all tested configurations fail. These observations contradict an unqualified interpretation of the earlier short-oligo RS2 pilot as evidence of local superiority. External performance depends on parity allocation, source-chunk count, screening and graph survival, not only on the codec name.')
P.para(d,'The external encoded density is 16,384/31,248 = 0.524322 payload bits per encoded nucleotide; the segmented local density is 16,384/31,680 = 0.517172. The external design uses 432 fewer bases, a 1.36% difference relative to the local budget. Density excludes primers and molecular copies for both methods. Source bytes per external oligo shrink with added parity, so more source chunks must be reconstructed from the same sixty-three droplets. Stronger per-droplet correction therefore trades against outer fountain overhead. Within this tested grid, the stronger inner code wins the strict-substitution file-level outcome up to 2%, but this does not prove a global optimum.')
P.h2(d,'Tradeoff, not reinstated superiority')
P.para(d,'These are different sequence constraints. RSNS enforces no adjacent repeated bases, whereas the clean-valid external run allows repeats up to five and screens GC fraction. A design allowing longer homopolymers is not equivalent under a synthesis rule that requires HP1. The local segmented control also assumes perfect order and chunk identity without charging index bases, while external droplets include four seed bytes. Known source length and codec settings are supplied to both. No physical oligo ordering errors, indels, strand dropout, primers, molecular-copy distribution or consensus-read budget are tested. There are still only four independent payload identities and one fixed deterministic external seed stream.')
P.para(d,'The defensible result is an observed software tradeoff and a warning about baseline parameterization: a pinned, clean-valid external port can outperform this local segmented design under strict substitutions at near-matched encoded-base budget, while accepting a different sequence constraint. It is not a claim that the local codec beats published storage systems, nor a claim that this port setting is universally best. Biological validation and an optimized constraint-matched frontier remain open. The initial local Fountain baseline, historical attempted-error rates, invalid-clean HP3 sweep and original short-oligo pilot remain archived separately rather than overwritten.')
P.para(d,'Reproduce per message: python3 experiments/fountain_valid_noise_audit.py --port /path/to/pinned/compiled/port --message N, N=0,1,2,3. The saved FASTA for each parity setting permits byte-for-byte encoded-sequence verification. Source: https://github.com/jdbrody/dna-fountain, commit f97c1b8a81f5c5b819209d5b5e26c9c8f4439495, linked by https://github.com/TeamErlich/dna-fountain. This software audit is distinct from the experimental paper at https://www.science.org/doi/10.1126/science.aaj2038.')
for row in d.tables[-1].rows:row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
for cell in d.tables[-1].rows[0].cells:
 for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
P.save(d,str(path));print('appended clean-valid noise audit')
